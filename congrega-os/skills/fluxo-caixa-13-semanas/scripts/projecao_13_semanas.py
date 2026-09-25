#!/usr/bin/env python3
"""Projeção de fluxo de caixa de 13 semanas pelo método direto (Congrega OS).

Lê um arquivo JSON de premissas (ver modelos/premissas_exemplo.json), gera os
eventos de caixa com datas ajustadas ao calendário bancário brasileiro, roda
os cenários base / otimista / pessimista e grava:

  <saida>/fluxo_13s_<cliente>_<inicio>.csv    semanal, todos os cenários
  <saida>/fluxo_13s_<cliente>_<inicio>.xlsx   (se openpyxl estiver instalado)
  <saida>/resumo_13s_<cliente>_<inicio>.md    diagnóstico e indicadores

Uso:
  python3 projecao_13_semanas.py premissas.json --saida ./saida
"""

import argparse
import calendar
import csv
import datetime as dt
import json
import re
import sys
from collections import defaultdict

SEMANAS = 13

CATEGORIAS_ENTRADA = [
    ("recebimento_carteira", "Recebimento da carteira (CR)"),
    ("vendas_novas", "Vendas novas"),
    ("outras_entradas", "Outras entradas operacionais"),
    ("captacoes", "Captações / aportes (financiamento)"),
]
CATEGORIAS_SAIDA = [
    ("folha", "Folha (salários, adiantamento, 13º, férias, rescisões)"),
    ("encargos_folha", "Encargos da folha (INSS, IRRF, FGTS)"),
    ("tributos", "Tributos (DAS, PIS/COFINS, IRPJ/CSLL, ICMS, ISS, CBS/IBS)"),
    ("fornecedores", "Fornecedores (CP)"),
    ("despesas_fixas", "Despesas fixas (aluguel, utilidades, sistemas)"),
    ("servico_divida", "Serviço da dívida (principal + juros)"),
    ("capex", "Investimentos (capex)"),
    ("distribuicoes", "Distribuição de lucros / pró-labore extra"),
    ("outras_saidas", "Outras saídas"),
]
CATEGORIAS_FINANCIAMENTO = {"captacoes", "servico_divida", "distribuicoes"}
CATEGORIAS_INVESTIMENTO = {"capex"}

CENARIOS_PADRAO = {
    "base": {"atraso_adicional_dias": 0, "fator_probabilidade": 1.0,
             "fator_vendas_novas": 1.0},
    "otimista": {"atraso_adicional_dias": -7, "fator_probabilidade": 1.05,
                 "fator_vendas_novas": 1.15},
    "pessimista": {"atraso_adicional_dias": 14, "fator_probabilidade": 0.85,
                   "fator_vendas_novas": 0.80},
}

PROBABILIDADE_PADRAO = {
    "a_vencer": 0.97,
    "1-30": 0.85,
    "31-60": 0.65,
    "61-90": 0.40,
    "91-180": 0.20,
    "180+": 0.05,
}


# --------------------------------------------------------------------------
# Calendário bancário
# --------------------------------------------------------------------------

def pascoa(ano):
    """Domingo de Páscoa (algoritmo de Meeus/Jones/Butcher)."""
    a = ano % 19
    b, c = divmod(ano, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l_ = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l_) // 451
    mes = (h + l_ - 7 * m + 114) // 31
    dia = ((h + l_ - 7 * m + 114) % 31) + 1
    return dt.date(ano, mes, dia)


def feriados_bancarios(ano):
    """Feriados nacionais (Lei 662/1949, Lei 6.802/1980, Lei 14.759/2023) e
    dias sem expediente bancário (carnaval e Corpus Christi)."""
    p = pascoa(ano)
    datas = {
        dt.date(ano, 1, 1),
        p - dt.timedelta(days=48),  # segunda de carnaval
        p - dt.timedelta(days=47),  # terça de carnaval
        p - dt.timedelta(days=2),   # sexta-feira santa
        dt.date(ano, 4, 21),
        dt.date(ano, 5, 1),
        p + dt.timedelta(days=60),  # Corpus Christi
        dt.date(ano, 9, 7),
        dt.date(ano, 10, 12),
        dt.date(ano, 11, 2),
        dt.date(ano, 11, 15),
        dt.date(ano, 12, 25),
    }
    if ano >= 2024:
        datas.add(dt.date(ano, 11, 20))
    return datas


class Calendario:
    def __init__(self, adicionais=()):
        self._cache = {}
        self._adicionais = {parse_data(d) for d in adicionais}

    def feriado(self, d):
        if d.year not in self._cache:
            self._cache[d.year] = feriados_bancarios(d.year)
        return d in self._cache[d.year] or d in self._adicionais

    def dia_util(self, d):
        return d.weekday() < 5 and not self.feriado(d)

    def ajustar(self, d, regra):
        if regra not in ("antecipa", "posterga"):
            return d
        passo = dt.timedelta(days=-1 if regra == "antecipa" else 1)
        while not self.dia_util(d):
            d += passo
        return d

    def ultimo_dia_util(self, ano, mes):
        d = dt.date(ano, mes, calendar.monthrange(ano, mes)[1])
        return self.ajustar(d, "antecipa")

    def quinto_dia_util_salario(self, ano, mes):
        """CLT art. 459, §1º. Conta sábado como dia útil (entendimento do MTE);
        se cair em sábado, o crédito precisa sair no dia bancário anterior."""
        d = dt.date(ano, mes, 1)
        contados = 0
        while True:
            if d.weekday() < 6 and not self.feriado(d):
                contados += 1
                if contados == 5:
                    return self.ajustar(d, "antecipa")
            d += dt.timedelta(days=1)


# --------------------------------------------------------------------------
# Utilitários
# --------------------------------------------------------------------------

def parse_data(txt):
    return dt.datetime.strptime(txt, "%Y-%m-%d").date()


def meses_da_janela(inicio, fim):
    ano, mes = inicio.year, inicio.month
    while (ano, mes) <= (fim.year, fim.month):
        yield ano, mes
        mes += 1
        if mes > 12:
            ano, mes = ano + 1, 1


def mes_anterior(ano, mes):
    return (ano - 1, 12) if mes == 1 else (ano, mes - 1)


def data_no_mes(cal, ano, mes, dia):
    if dia == "ultimo_dia_util":
        return cal.ultimo_dia_util(ano, mes)
    if dia == "quinto_dia_util":
        return cal.quinto_dia_util_salario(ano, mes)
    ultimo = calendar.monthrange(ano, mes)[1]
    return dt.date(ano, mes, min(int(dia), ultimo))


def faixa_atraso(dias):
    if dias <= 0:
        return "a_vencer"
    if dias <= 30:
        return "1-30"
    if dias <= 60:
        return "31-60"
    if dias <= 90:
        return "61-90"
    if dias <= 180:
        return "91-180"
    return "180+"


def brl(v):
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-R$ {s}" if v < 0 else f"R$ {s}"


def slug(txt):
    return re.sub(r"[^a-z0-9]+", "-", txt.lower()).strip("-") or "cliente"


# --------------------------------------------------------------------------
# Geração de eventos
# --------------------------------------------------------------------------

def evento(data, valor, categoria, descricao):
    return {"data": data, "valor": float(valor), "categoria": categoria,
            "descricao": descricao}


def eventos_carteira(p, cenario, inicio, fim, avisos):
    prob_tab = {**PROBABILIDADE_PADRAO, **p.get("probabilidade_recebimento", {})}
    atraso_medio = p.get("atraso_medio_dias", 0) + cenario["atraso_adicional_dias"]
    atraso_vencidos = p.get("prazo_cobranca_vencidos_dias", 14)
    fator = cenario["fator_probabilidade"]
    cal = p["_cal"]
    saida, fora = [], 0.0
    for t in p.get("contas_receber", []):
        venc = parse_data(t["vencimento"])
        dias_atraso = (inicio - venc).days
        prob = t.get("probabilidade", prob_tab[faixa_atraso(dias_atraso)])
        prob = max(0.0, min(1.0, prob * fator))
        if dias_atraso > 0:
            data = inicio + dt.timedelta(days=max(0, atraso_vencidos + cenario["atraso_adicional_dias"]))
        else:
            data = venc + dt.timedelta(days=atraso_medio)
            data = max(data, inicio)
        data = cal.ajustar(data, "posterga")
        liquido = t["valor"] - t.get("retencoes", 0.0)
        esperado = liquido * prob
        if data > fim:
            fora += esperado
            continue
        saida.append(evento(data, esperado, "recebimento_carteira",
                            f"{t.get('cliente', 'CR')} venc. {venc:%d/%m} ({prob:.0%})"))
    if fora:
        avisos.append(f"{brl(fora)} esperados da carteira caem fora das 13 semanas.")
    return saida


def distribuir_semanal(p, semanas, chave, fator, categoria, descricao, sinal):
    valores = p.get(chave, [])
    if valores and len(valores) != SEMANAS:
        raise ValueError(f"{chave} precisa ter 13 valores.")
    cal = p["_cal"]
    saida = []
    for (ini, fim), valor in zip(semanas, valores):
        dias = [ini + dt.timedelta(days=i) for i in range((fim - ini).days + 1)]
        uteis = [d for d in dias if cal.dia_util(d)] or [cal.ajustar(fim, "posterga")]
        parcela = sinal * abs(valor) * fator / len(uteis)
        saida += [evento(d, parcela, categoria, descricao) for d in uteis]
    return saida


def eventos_vendas_novas(p, cenario, semanas):
    return distribuir_semanal(p, semanas, "vendas_novas_semanais", cenario["fator_vendas_novas"],
                              "vendas_novas", "Vendas novas (pipeline ponderado)", 1)


def eventos_compras_novas(p, cenario, semanas):
    # Compras acompanham o volume de vendas, salvo se o cenário disser o contrário
    fator = cenario.get("fator_compras_novas", cenario["fator_vendas_novas"])
    return distribuir_semanal(p, semanas, "compras_novas_semanais", fator,
                              "fornecedores", "Compras novas (a prazo, vencendo na janela)", -1)


def eventos_contas_pagar(p, inicio):
    cal = p["_cal"]
    saida = []
    for t in p.get("contas_pagar", []):
        venc = parse_data(t["vencimento"])
        data = cal.ajustar(max(venc, inicio), "posterga")
        cat = t.get("categoria", "fornecedores")
        saida.append(evento(data, -abs(t["valor"]), cat,
                            f"{t.get('fornecedor', 'CP')} venc. {venc:%d/%m}"))
    return saida


def eventos_folha(p, inicio, fim):
    f = p.get("folha")
    if not f:
        return []
    cal = p["_cal"]
    bruto = f.get("salario_bruto_mensal", 0.0)
    liquido = f["salarios_liquidos_mensal"]
    adi = f.get("adiantamento", {})
    pct_adi = adi.get("percentual", 0.0)
    saida = []
    for ano, mes in meses_da_janela(inicio, fim):
        ref_ano, ref_mes = mes_anterior(ano, mes)
        # Saldo do salário do mês anterior: até o 5º dia útil (CLT art. 459, §1º)
        saida.append(evento(cal.quinto_dia_util_salario(ano, mes),
                            -liquido * (1 - pct_adi), "folha",
                            f"Salários {ref_mes:02d}/{ref_ano} (saldo)"))
        if pct_adi:
            saida.append(evento(cal.ajustar(data_no_mes(cal, ano, mes, adi.get("dia", 20)), "antecipa"),
                                -liquido * pct_adi, "folha",
                                f"Adiantamento salarial {mes:02d}/{ano}"))
        # INSS (Lei 8.212/1991, art. 30, I, b) e IRRF (Lei 11.196/2005, art. 70, I, d): dia 20, antecipa
        if f.get("inss_irrf_mensal"):
            saida.append(evento(cal.ajustar(dt.date(ano, mes, 20), "antecipa"),
                                -f["inss_irrf_mensal"], "encargos_folha",
                                f"INSS + IRRF folha {ref_mes:02d}/{ref_ano} (DCTFWeb)"))
        # FGTS Digital: dia 20 (Lei 8.036/1990, art. 15, red. Lei 14.438/2022), antecipa
        if f.get("fgts_mensal"):
            saida.append(evento(cal.ajustar(dt.date(ano, mes, 20), "antecipa"),
                                -f["fgts_mensal"], "encargos_folha",
                                f"FGTS {ref_mes:02d}/{ref_ano}"))
        if f.get("decimo_terceiro", True) and bruto:
            fgts_aliq = f.get("fgts_mensal", 0.0) / bruto if bruto else 0.08
            if mes == 11:
                # 1ª parcela até 30/11 (Lei 4.749/1965, art. 2º)
                saida.append(evento(cal.ajustar(dt.date(ano, 11, 30), "antecipa"),
                                    -bruto * 0.5, "folha", "13º salário - 1ª parcela"))
            if mes == 12:
                # 2ª parcela até 20/12 (Lei 4.749/1965, art. 1º), líquida de descontos
                descontos = max(0.0, bruto - liquido)
                saida.append(evento(cal.ajustar(dt.date(ano, 12, 20), "antecipa"),
                                    -(bruto * 0.5 - descontos), "folha",
                                    "13º salário - 2ª parcela (estimada)"))
                saida.append(evento(cal.ajustar(dt.date(ano, 12, 20), "antecipa"),
                                    -f.get("inss_irrf_mensal", 0.0), "encargos_folha",
                                    "INSS + IRRF sobre 13º (estimado)"))
                saida.append(evento(cal.ajustar(dt.date(ano, 12, 20), "antecipa"),
                                    -bruto * 0.5 * fgts_aliq, "encargos_folha",
                                    "FGTS sobre 1ª parcela do 13º"))
            if mes == 1:
                saida.append(evento(cal.ajustar(dt.date(ano, 1, 20), "antecipa"),
                                    -bruto * 0.5 * fgts_aliq, "encargos_folha",
                                    "FGTS sobre 2ª parcela do 13º"))
    return saida


def eventos_recorrentes(p, inicio, fim):
    cal = p["_cal"]
    saida = []
    for r in p.get("recorrentes", []):
        meses = r.get("meses")
        sinal = 1 if r.get("tipo", "saida") == "entrada" else -1
        for ano, mes in meses_da_janela(inicio, fim):
            if meses and mes not in meses:
                continue
            data = cal.ajustar(data_no_mes(cal, ano, mes, r["dia"]), r.get("ajuste", "antecipa"))
            saida.append(evento(data, sinal * abs(r["valor"]), r.get("categoria", "outras_saidas"),
                                f"{r['descricao']} {mes:02d}/{ano}"))
    return saida


def eventos_avulsos(lista, p):
    cal = p["_cal"]
    saida = []
    for e in lista:
        sinal = 1 if e.get("tipo", "saida") == "entrada" else -1
        cat = e.get("categoria", "outras_entradas" if sinal > 0 else "outras_saidas")
        data = cal.ajustar(parse_data(e["data"]), e.get("ajuste", "antecipa"))
        saida.append(evento(data, sinal * abs(e["valor"]), cat, e["descricao"]))
    return saida


# --------------------------------------------------------------------------
# Projeção
# --------------------------------------------------------------------------

def janelas(inicio):
    return [(inicio + dt.timedelta(days=7 * i), inicio + dt.timedelta(days=7 * i + 6))
            for i in range(SEMANAS)]


def projetar(p, nome, cenario):
    inicio = parse_data(p["data_inicio"])
    semanas = janelas(inicio)
    fim = semanas[-1][1]
    avisos = []
    eventos = []
    eventos += eventos_carteira(p, cenario, inicio, fim, avisos)
    eventos += eventos_vendas_novas(p, cenario, semanas)
    eventos += eventos_compras_novas(p, cenario, semanas)
    eventos += eventos_contas_pagar(p, inicio)
    eventos += eventos_folha(p, inicio, fim)
    eventos += eventos_recorrentes(p, inicio, fim)
    eventos += eventos_avulsos(p.get("eventos", []), p)
    eventos += eventos_avulsos(cenario.get("eventos_extras", []), p)
    eventos = [e for e in eventos if inicio <= e["data"] <= fim and abs(e["valor"]) >= 0.005]

    saldo = p["saldo_inicial"] - p.get("caixa_restrito", 0.0)
    minimo = p.get("saldo_minimo", 0.0)
    limite = p.get("limite_credito_disponivel", 0.0)
    linhas = []
    for n, (ini, fim_s) in enumerate(semanas, start=1):
        cats = defaultdict(float)
        for e in eventos:
            if ini <= e["data"] <= fim_s:
                cats[e["categoria"]] += e["valor"]
        entradas = sum(v for v in cats.values() if v > 0)
        saidas = sum(v for v in cats.values() if v < 0)
        fin = sum(v for k, v in cats.items() if k in CATEGORIAS_FINANCIAMENTO)
        inv = sum(v for k, v in cats.items() if k in CATEGORIAS_INVESTIMENTO)
        oper = entradas + saidas - fin - inv
        abertura = saldo
        saldo += entradas + saidas
        linhas.append({
            "cenario": nome, "semana": n, "inicio": ini.isoformat(), "fim": fim_s.isoformat(),
            "saldo_inicial": round(abertura, 2),
            **{k: round(cats.get(k, 0.0), 2) for k, _ in CATEGORIAS_ENTRADA + CATEGORIAS_SAIDA},
            "total_entradas": round(entradas, 2), "total_saidas": round(saidas, 2),
            "fluxo_operacional": round(oper, 2), "fluxo_investimento": round(inv, 2),
            "fluxo_financiamento": round(fin, 2), "fluxo_liquido": round(entradas + saidas, 2),
            "saldo_final": round(saldo, 2), "saldo_com_limite": round(saldo + limite, 2),
            "alerta": "ABAIXO DO MINIMO" if saldo < minimo else "",
        })
    diario = []
    ini1, fim1 = semanas[0]
    saldo_d = p["saldo_inicial"] - p.get("caixa_restrito", 0.0)
    d = ini1
    while d <= fim1:
        do_dia = sorted((e for e in eventos if e["data"] == d), key=lambda e: -e["valor"])
        for e in do_dia:
            saldo_d += e["valor"]
            diario.append({"data": d.isoformat(), "descricao": e["descricao"],
                           "categoria": e["categoria"], "valor": round(e["valor"], 2),
                           "saldo": round(saldo_d, 2)})
        d += dt.timedelta(days=1)
    return linhas, diario, avisos, eventos


def diagnostico(linhas, p):
    minimo = p.get("saldo_minimo", 0.0)
    limite = p.get("limite_credito_disponivel", 0.0)
    pior = min(linhas, key=lambda l: l["saldo_final"])
    abaixo = [l["semana"] for l in linhas if l["saldo_final"] < minimo]
    estouro = [l["semana"] for l in linhas if l["saldo_com_limite"] < 0]
    negativo = [l["semana"] for l in linhas if l["saldo_final"] < 0]
    return {
        "menor_saldo": pior["saldo_final"], "semana_menor_saldo": pior["semana"],
        "semanas_abaixo_minimo": abaixo, "semana_de_perigo": abaixo[0] if abaixo else None,
        "semana_caixa_negativo": negativo[0] if negativo else None,
        "semanas_estouro_limite": estouro,
        "necessidade_caixa": max(0.0, minimo - pior["saldo_final"]),
        "saldo_final": linhas[-1]["saldo_final"],
        "limite": limite,
    }


def indicadores_capital_giro(p):
    ind = p.get("indicadores")
    if not ind:
        return None
    dias = ind.get("dias_base", 360)
    r = {}
    if ind.get("receita_bruta_anual") and "contas_receber" in ind:
        r["PMR"] = ind["contas_receber"] / ind["receita_bruta_anual"] * dias
    if ind.get("cmv_anual") and "estoques" in ind:
        r["PME"] = ind["estoques"] / ind["cmv_anual"] * dias
    base_pmp = ind.get("compras_anuais") or ind.get("cmv_anual")
    if base_pmp and "fornecedores" in ind:
        r["PMP"] = ind["fornecedores"] / base_pmp * dias
    if "PMR" in r:
        r["ciclo_operacional"] = r["PMR"] + r.get("PME", 0.0)
        if "PMP" in r:
            r["ciclo_financeiro"] = r["ciclo_operacional"] - r["PMP"]
    if "ativo_circulante_operacional" in ind and "passivo_circulante_operacional" in ind:
        r["NCG"] = ind["ativo_circulante_operacional"] - ind["passivo_circulante_operacional"]
        if ind.get("receita_bruta_anual"):
            r["NCG_dias_receita"] = r["NCG"] / ind["receita_bruta_anual"] * dias
    return r


# --------------------------------------------------------------------------
# Saídas
# --------------------------------------------------------------------------

def gravar_csv(caminho, linhas):
    with open(caminho, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(linhas[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(linhas)


def gravar_xlsx(caminho, p, resultados, diario_base, cg):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    navy = PatternFill("solid", fgColor="1F2A44")
    alerta = PatternFill("solid", fgColor="F8D7DA")
    branco = Font(color="FFFFFF", bold=True)
    fmt = '#,##0.00;[Red]-#,##0.00'
    wb = Workbook()
    wb.remove(wb.active)

    for nome, (linhas, _, _, _) in resultados.items():
        ws = wb.create_sheet(nome.capitalize())
        ws.append(["Linha"] + [f"S{l['semana']}\n{l['inicio'][8:10]}/{l['inicio'][5:7]}" for l in linhas]
                  + ["Total 13s"])
        blocos = ([("saldo_inicial", "Saldo inicial disponível")]
                  + [(None, "ENTRADAS")] + CATEGORIAS_ENTRADA + [("total_entradas", "Total de entradas")]
                  + [(None, "SAÍDAS")] + CATEGORIAS_SAIDA + [("total_saidas", "Total de saídas")]
                  + [(None, "RESULTADO")]
                  + [("fluxo_operacional", "Fluxo operacional"), ("fluxo_investimento", "Fluxo de investimento"),
                     ("fluxo_financiamento", "Fluxo de financiamento"), ("fluxo_liquido", "Fluxo líquido"),
                     ("saldo_final", "Saldo final"), ("saldo_com_limite", "Saldo + limite de crédito")])
        for chave, rotulo in blocos:
            if chave is None:
                ws.append([rotulo])
                ws.cell(ws.max_row, 1).font = Font(bold=True)
                continue
            valores = [l[chave] for l in linhas]
            total = "" if chave in ("saldo_inicial", "saldo_final", "saldo_com_limite") else sum(valores)
            ws.append([rotulo] + valores + [total])
            if chave.startswith(("total", "saldo_final", "fluxo_liquido")):
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
        for c in ws[1]:
            c.fill, c.font = navy, branco
            c.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.row_dimensions[1].height = 32
        linha_saldo = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Saldo final")
        for i, l in enumerate(linhas, start=2):
            if l["alerta"]:
                ws.cell(linha_saldo, i).fill = alerta
        for row in ws.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.number_format = fmt
        ws.column_dimensions["A"].width = 52
        for col in range(2, SEMANAS + 3):
            ws.column_dimensions[get_column_letter(col)].width = 14
        ws.freeze_panes = "B2"

    ws = wb.create_sheet("Semana 1 diária")
    ws.append(["Data", "Descrição", "Categoria", "Valor", "Saldo"])
    for r in diario_base:
        ws.append([r["data"], r["descricao"], r["categoria"], r["valor"], r["saldo"]])
    for c in ws[1]:
        c.fill, c.font = navy, branco
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = fmt
    for col, w in zip("ABCDE", (12, 52, 24, 16, 16)):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Resumo", 0)
    ws.append(["Cliente", p["cliente"]])
    ws.append(["Início da projeção", p["data_inicio"]])
    ws.append(["Saldo mínimo de segurança", p.get("saldo_minimo", 0.0)])
    ws.append(["Limite de crédito disponível", p.get("limite_credito_disponivel", 0.0)])
    ws.append([])
    ws.append(["Cenário", "Menor saldo", "Semana", "Semana de perigo", "Necessidade de caixa", "Saldo final S13"])
    for c in ws[ws.max_row]:
        c.fill, c.font = navy, branco
    for nome, (linhas, _, _, _) in resultados.items():
        d = diagnostico(linhas, p)
        ws.append([nome, d["menor_saldo"], d["semana_menor_saldo"], d["semana_de_perigo"] or "-",
                   d["necessidade_caixa"], d["saldo_final"]])
    if cg:
        ws.append([])
        ws.append(["Indicador de capital de giro", "Valor"])
        for c in ws[ws.max_row]:
            c.fill, c.font = navy, branco
        for k, v in cg.items():
            ws.append([k, round(v, 2)])
    for row in ws.iter_rows(min_row=3, min_col=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = fmt
    ws.column_dimensions["A"].width = 32
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20
    wb.save(caminho)
    return True


def gravar_resumo(caminho, p, resultados, cg):
    out = [f"# Fluxo de caixa 13 semanas - {p['cliente']}", "",
           f"Início: {p['data_inicio']} | Saldo inicial disponível: "
           f"{brl(p['saldo_inicial'] - p.get('caixa_restrito', 0.0))} | "
           f"Saldo mínimo: {brl(p.get('saldo_minimo', 0.0))} | "
           f"Limite de crédito: {brl(p.get('limite_credito_disponivel', 0.0))}", "",
           "| Cenário | Menor saldo | Semana | 1ª semana abaixo do mínimo | Semanas abaixo do mínimo | Caixa negativo em | Necessidade de caixa | Saldo S13 |",
           "|---|---|---|---|---|---|---|---|"]
    avisos_todos = []
    for nome, (linhas, _, avisos, _) in resultados.items():
        d = diagnostico(linhas, p)
        sem = lambda n: f"S{n}" if n else "-"
        out.append(f"| {nome} | {brl(d['menor_saldo'])} | S{d['semana_menor_saldo']} | "
                   f"{sem(d['semana_de_perigo'])} | {len(d['semanas_abaixo_minimo'])} | "
                   f"{sem(d['semana_caixa_negativo'])} | "
                   f"{brl(d['necessidade_caixa'])} | {brl(d['saldo_final'])} |")
        if d["semanas_estouro_limite"]:
            avisos_todos.append(f"[{nome}] saldo ultrapassa o limite de crédito nas semanas "
                                f"{', '.join('S' + str(s) for s in d['semanas_estouro_limite'])}.")
        avisos_todos += [f"[{nome}] {a}" for a in avisos]
    if cg:
        out += ["", "## Capital de giro", ""]
        out += [f"- {k}: {v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") for k, v in cg.items()]
    if avisos_todos:
        out += ["", "## Avisos", ""] + [f"- {a}" for a in avisos_todos]
    with open(caminho, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out) + "\n")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("premissas", help="arquivo JSON de premissas")
    ap.add_argument("--saida", default=".", help="diretório de saída")
    args = ap.parse_args()

    with open(args.premissas, encoding="utf-8") as fh:
        p = json.load(fh)
    p["_cal"] = Calendario(p.get("feriados_adicionais", []))

    cenarios = {k: {**CENARIOS_PADRAO.get(k, CENARIOS_PADRAO["base"]), **v}
                for k, v in {**{k: {} for k in CENARIOS_PADRAO}, **p.get("cenarios", {})}.items()}
    resultados = {nome: projetar(p, nome, c) for nome, c in cenarios.items()}
    cg = indicadores_capital_giro(p)

    base = f"{slug(p['cliente'])}_{p['data_inicio']}"
    todas = [l for linhas, _, _, _ in resultados.values() for l in linhas]
    gravar_csv(f"{args.saida}/fluxo_13s_{base}.csv", todas)
    tem_xlsx = gravar_xlsx(f"{args.saida}/fluxo_13s_{base}.xlsx", p, resultados,
                           resultados.get("base", next(iter(resultados.values())))[1], cg)
    print(gravar_resumo(f"{args.saida}/resumo_13s_{base}.md", p, resultados, cg))
    print()
    print(f"Arquivos gravados em {args.saida}: CSV, resumo .md" + (", XLSX" if tem_xlsx else
          " (XLSX não gerado: instale openpyxl)"))


if __name__ == "__main__":
    sys.exit(main())
